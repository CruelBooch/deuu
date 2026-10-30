"""生成申请书的 4 个附件（A4 docx）"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

DOCS = r"D:\workspace\deu\docs"
SHOTS = os.path.join(DOCS, "shots")
ROOT = r"D:\workspace\deu"


def new_doc():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2.2)
    s.left_margin, s.right_margin = Cm(2.6), Cm(2.6)
    n = d.styles["Normal"]
    n.font.name = "宋体"
    n.font.size = Pt(12)
    n.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return d


def rf(run, name="宋体", size=12, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def p(d, text="", *, indent=True, size=12, font="宋体", bold=False,
      align=None, sa=4, line=1.35, sb=0):
    para = d.add_paragraph()
    pf = para.paragraph_format
    pf.first_line_indent = Pt(size * 2) if indent else Pt(0)
    pf.space_after, pf.space_before, pf.line_spacing = Pt(sa), Pt(sb), line
    if align is not None:
        para.alignment = align
    if text:
        rf(para.add_run(text), font, size, bold)
    return para


def h(d, text, level=1):
    return p(d, text, indent=False, font="黑体", bold=True,
             size=15 if level == 1 else 12, sb=0 if level == 1 else 10, sa=8 if level == 1 else 4)


def bullet(d, text, size=12, sa=2):
    para = d.add_paragraph()
    pf = para.paragraph_format
    pf.left_indent, pf.first_line_indent = Pt(24), Pt(-12)
    pf.space_after, pf.line_spacing = Pt(sa), 1.35
    rf(para.add_run("· " + text), "宋体", size)


def table(d, rows, widths, size=10.5):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.autofit = False
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = Cm(widths[j])
            cp = c.paragraphs[0]
            cp.paragraph_format.space_after = Pt(2)
            cp.paragraph_format.space_before = Pt(2)
            cp.paragraph_format.line_spacing = 1.1
            rf(cp.add_run(val), "宋体", size, bold=(i == 0))
    return t


# ============ 附件一：功能截图 ============
d = new_doc()
h(d, "附件一：应用功能截图")
p(d, "以下截图取自 Android 版实机运行效果；Windows 版界面完全一致，仅窗口标题栏不同。"
     "截图环境为深色模式，应用会跟随系统自动切换明暗主题。", size=9.5, sa=6)

items = [
    ("01-home.png", "句子翻译：原文、译文与语法讲解"),
    ("02-dict.png", "查词结果：词性、性、复数、音标、释义"),
    ("03-suggest.png", "输入时给出离线词库候选"),
    ("04-history.png", "历史记录：按时间分组、可单条删除"),
    ("05-vocab.png", "生词本：收藏、自定义分组、批量操作"),
    ("06-settings.png", "设置：API Key、接口地址、模型、主题"),
]

t = d.add_table(rows=2, cols=3)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for idx, (fname, cap) in enumerate(items):
    r, c = idx // 3, idx % 3
    cell = t.cell(r, c)
    cell.width = Cm(5.1)
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(2)
    path = os.path.join(SHOTS, fname)
    if os.path.exists(path):
        cp.add_run().add_picture(path, width=Cm(4.2))
    cap_p = cell.add_paragraph()
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_p.paragraph_format.space_after = Pt(6)
    rf(cap_p.add_run(cap), "宋体", 8)

p(d, "说明：离线词典内置约 2 万条德语词条，断网状态下即可查词；"
     "句子翻译与语法讲解需要用户自行配置 AI 接口，未配置时不影响查词与生词本功能。",
  size=10, indent=False, sa=0, sb=8)
d.save(os.path.join(DOCS, "附件1-应用功能截图.docx"))

# ============ 附件二：源代码与技术文档清单 ============
src_files, src_lines = 0, 0
for base, _, files in os.walk(os.path.join(ROOT, "src")):
    for f in files:
        if f.endswith((".ts", ".tsx", ".css", ".html")):
            src_files += 1
            with open(os.path.join(base, f), encoding="utf-8", errors="ignore") as fh:
                src_lines += sum(1 for _ in fh)

d = new_doc()
h(d, "附件二：源代码与技术文档清单")
p(d, "本应用为本人独立开发的桌面与移动端应用，源代码完整可交付学校备份与审查。")

h(d, "一、基本信息", 2)
table(d, [
    ["项目", "内容"],
    ["应用名称", "deuu（德语学习辅助工具）"],
    ["版本", "0.1.0"],
    ["开发语言", "TypeScript"],
    ["界面框架", "React 18"],
    ["桌面外壳", "Electron 33（Windows）"],
    ["移动外壳", "Capacitor 6（Android）"],
    ["源代码规模", f"约 {src_files} 个源文件、{src_lines} 行代码"],
    ["构建产物", "Windows 安装包 / 便携版、Android APK、Web 版"],
], [3.2, 12.0])

h(d, "二、代码结构", 2)
for line in [
    "src/core/          纯业务逻辑（AI 解析、词典索引），不依赖任何平台 API",
    "src/main/          桌面端主进程：窗口、设置、历史、生词本、词典加载",
    "src/preload/       桌面端安全桥接（window.api）",
    "src/renderer/      React 界面（桌面与移动端共用同一套）",
    "src/renderer/src/platform/  平台适配层（Capacitor / 未来的 iOS、鸿蒙）",
    "dict/              离线词典数据（约 2 万词条）",
    "scripts/           词典构建、构建环境与打包脚本",
]:
    bullet(d, line, size=10.5, sa=1)

h(d, "三、技术文档", 2)
table(d, [
    ["文档", "内容"],
    ["README.md", "项目说明、开发与打包流程、已知问题"],
    ["docs/PLATFORM.md", "跨平台架构说明、平台能力契约、新增平台的方法"],
    ["docs/ANDROID-PLAN.md", "Android 端方案与实施记录"],
    ["scripts/build_dict.py", "离线词典构建脚本（从维基词典 dump 解析）"],
], [5.0, 10.2])

h(d, "四、交付与审查方式", 2)
p(d, "本人可随时提供以下任一种方式，供学校审查：")
bullet(d, "现场演示：携带设备，演示查词、翻译、生词本等全部功能；")
bullet(d, "源代码压缩包：完整工程（不含第三方依赖），可在校内电脑上构建；")
bullet(d, "构建说明：提供从零构建可执行文件的完整步骤文档。")

p(d, "本人承诺：源代码中不包含任何恶意代码、后门、数据回传逻辑；"
     "所有联网行为仅限用户主动发起的翻译请求。", sa=2, sb=6)
d.save(os.path.join(DOCS, "附件2-源代码与技术文档清单.docx"))

# ============ 附件三：隐私说明 ============
d = new_doc()
h(d, "附件三：隐私说明")
p(d, "本说明适用于 deuu（德语学习辅助工具，以下简称“本应用”），"
     "用于向学校说明本应用的数据处理方式。")

h(d, "一、我们不收集任何个人信息", 2)
p(d, "本应用没有账号体系，不需要注册、登录，不收集姓名、学号、手机号、邮箱、"
     "地理位置、通讯录、相册等任何个人身份信息。")

h(d, "二、数据保存在哪里", 2)
p(d, "本应用的设置、查询历史、生词本全部保存在用户本机（手机或电脑的本地存储中），"
     "不会上传到任何服务器，也不会有任何云端备份。"
     "卸载应用或清除数据后，这些内容即被删除。")

h(d, "三、关于 AI 翻译功能", 2)
p(d, "单词查询完全离线运行，不产生任何网络请求。"
     "仅当用户主动使用“句子翻译”或“AI 详细查询”时，本应用才会把用户输入的那段文字，"
     "发送到用户自己在设置中配置的第三方 AI 接口（如 DeepSeek、OpenAI 等）。"
     "接口地址与密钥均由用户自行填写，开发者不接触、不保存这些内容。")
p(d, "本应用不向第三方发送任何身份信息，也不在后台自动发起请求。")

h(d, "四、权限说明", 2)
p(d, "本应用仅申请“网络访问”一项权限（用于翻译功能），"
     "不申请相机、麦克风、定位、通讯录、存储读写等敏感权限。")

h(d, "五、内容说明", 2)
p(d, "本应用为德语学习工具，不含社交、评论、广告、付费内容，"
     "不存在用户生成内容（UGC）功能。")

h(d, "六、联系方式", 2)
p(d, "如对本说明有任何疑问，可联系开发者：")
p(d, "姓名：×××　　班级：×× 年级 ×× 班　　邮箱：______________", indent=False, sa=0)
d.save(os.path.join(DOCS, "附件3-隐私说明.docx"))

# ============ 附件四：会费豁免政策说明 ============
d = new_doc()
h(d, "附件四：Apple 会费豁免政策说明")
p(d, "以下为苹果公司官方《Apple Developer Program 会费豁免》政策的要点摘录，"
     "供学校审批时参考。完整内容请以苹果公司官网页面为准。")

h(d, "一、政策来源", 2)
table(d, [
    ["项目", "内容"],
    ["政策名称", "Apple Developer Program 会费豁免（Fee Waivers）"],
    ["发布方", "Apple Inc."],
    ["官方页面", "developer.apple.com/cn/help/account/membership/fee-waivers/"],
    ["查询日期", "2026 年 ×× 月 ×× 日"],
], [3.2, 12.0])

h(d, "二、适用对象", 2)
p(d, "根据该政策，以下三类组织可申请免除 Apple Developer Program 的年度会费：")
bullet(d, "非营利组织（Nonprofit organizations）")
bullet(d, "经认证的教育机构（Accredited educational institutions）")
bullet(d, "政府机构（Government entities）")

h(d, "三、对本校的意义", 2)
p(d, "学校作为经认证的教育机构，符合上述适用条件。"
     "以学校名义注册组织账号并提交会费豁免申请后，即可免除每年 99 美元的会费，"
     "学校无需为此支出任何费用。")

h(d, "四、申请通常需要的材料", 2)
table(d, [
    ["材料", "说明"],
    ["D-U-N-S 编码", "邓白氏企业编码，可免费申请，通常 1–2 周出具"],
    ["机构证明", "事业单位法人证书 / 办学许可证等"],
    ["机构邮箱与联系人", "用于苹果公司身份核验"],
    ["豁免申请表", "在 Apple 开发者官网提交"],
], [4.0, 11.2])

h(d, "五、注意事项", 2)
p(d, "本附件仅为政策摘录说明，不构成任何承诺。是否申请、如何申请，"
     "请以苹果公司官网的最新要求及学校相关部门的判断为准。"
     "建议打印官网页面作为正式附件一并留存。")
d.save(os.path.join(DOCS, "附件4-会费豁免政策说明.docx"))

print("attachments done")
print("src files:", src_files, "lines:", src_lines)

